import csv

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Конвертирует CSV в файл Excel (XLSX).

    Первая строка CSV попадает в таблицу как обычная строка-заголовок:
    отдельного оформления шапки не делается. Лист называется "Sheet1".

    Ширина каждой колонки подбирается по самому длинному значению в ней,
    но не меньше 8 символов. Если в какой-то строке колонок меньше, чем
    в заголовке, недостающие ячейки просто не учитываются при подсчёте.

    Значения переносятся как есть, строками — CSV не хранит типы, поэтому
    числа в Excel тоже окажутся текстом.

    Args:
        csv_path: путь к входному CSV в кодировке UTF-8.
        xlsx_path: путь к создаваемому XLSX; существующий перезаписывается.

    Returns:
        None.

    Raises:
        FileNotFoundError: если входного файла нет.
        ValueError: если CSV пустой.
    """
    # reader, а не DictReader: в Excel заголовок — просто первая строка
    with open(csv_path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
        if not rows:
            raise ValueError("Пустой CSV")

    wb = Workbook()
    ws = wb.active
    # по умолчанию лист зовётся "Sheet", а по заданию нужен "Sheet1"
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)

    # второй цикл, не вложенный в первый: сначала заполнили, теперь меряем
    for i in range(len(rows[0])):
        # if i < len(row) — страховка от рядов короче заголовка
        longest = max(len(row[i]) for row in rows if i < len(row))
        # get_column_letter: Excel адресует колонки буквами, нумерация с 1
        ws.column_dimensions[get_column_letter(i + 1)].width = max(longest, 8)

    # сохраняем один раз, после всех правок
    wb.save(xlsx_path)


if __name__ == "__main__":
    from openpyxl import load_workbook

    # пути относительные, запускать из корня проекта
    pairs = [
        ("data/samples/cities.csv", "data/out/cities.xlsx"),
        ("data/samples/people.csv", "data/out/people.xlsx"),
    ]
    for src, dst in pairs:
        print("=" * 52)
        print(f"CSV -> XLSX: {src}")
        print("=" * 52)
        print(f"--- {src}")
        with open(src, encoding="utf-8") as f:
            print(f.read().rstrip())

        csv_to_xlsx(src, dst)

        # читаем готовый файл обратно — доказательство, что он не битый
        ws = load_workbook(dst).active
        print(f"--- {dst}")
        print(f"лист: {ws.title}, строк: {ws.max_row}, колонок: {ws.max_column}")
        for row in ws.iter_rows(values_only=True):
            print("   ", row)
        widths = {k: round(v.width) for k, v in ws.column_dimensions.items()}
        print(f"ширины колонок: {widths}  (минимум 8)")
        print()
